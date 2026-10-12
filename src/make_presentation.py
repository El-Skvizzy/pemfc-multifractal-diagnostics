# -*- coding: utf-8 -*-
"""
make_presentation.py — сборка презентации конкурсной работы (10 слайдов)
по инструкции docs/Оформление_научной_статьи/Инструкция_по_презентации.md.

Запуск: python src/make_presentation.py ["$PEMFC_PROJECT_ROOT"]
Выход: docs/Кондрашов - Цифровые технологии - презентация.pptx
"""
import os
import re
import sys

from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

sys.stdout.reconfigure(encoding='utf-8')

FONT = 'Times New Roman'
BLACK = RGBColor(0x1A, 0x1A, 0x1A)
DARK = RGBColor(0x1F, 0x4E, 0x79)
GRAY = RGBColor(0x80, 0x80, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xEA, 0xF1, 0xF8)

PROJECT_ROOT = os.environ.get("PEMFC_PROJECT_ROOT") or os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), os.pardir))
ROOT = sys.argv[1] if len(sys.argv) > 1 else PROJECT_ROOT
FIG = os.path.join(ROOT, 'results', 'figures')
if not os.path.isdir(FIG):
    FIG = os.path.join(ROOT, 'figures')
OUT = os.path.join(ROOT, 'docs', 'Кондрашов - Цифровые технологии - презентация.pptx')

SW, SH = Inches(13.333), Inches(7.5)

INLINE = re.compile(r'(\*\*.+?\*\*|\^[^^]+?\^|~[^~]+?~)')


def _set_run(r, text, size, bold=False, italic=False, color=BLACK):
    r.text = text
    f = r.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color


def add_rich(p, text, size=20, bold=False, italic=False, color=BLACK):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            add_rich(p, part[2:-2], size, True, italic, color)
        elif part.startswith('^') and part.endswith('^') and len(part) > 2:
            r = p.add_run(); _set_run(r, part[1:-1], size, bold, italic, color)
            r._r.get_or_add_rPr().set('baseline', '30000')
        elif part.startswith('~') and part.endswith('~') and len(part) > 2:
            r = p.add_run(); _set_run(r, part[1:-1], size, bold, italic, color)
            r._r.get_or_add_rPr().set('baseline', '-25000')
        else:
            r = p.add_run(); _set_run(r, part, size, bold, italic, color)


def add_textbox(slide, l, t, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    return tb, tf


def title_bar(slide, title, num=None):
    tb, tf = add_textbox(slide, 0.6, 0.35, 12.1, 0.9)
    _set_run(tf.paragraphs[0].add_run(), title, 28, bold=True)
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.18),
                                  Inches(2.4), Pt(2.6))
    rule.fill.solid(); rule.fill.fore_color.rgb = DARK
    rule.line.fill.background()
    rule.shadow.inherit = False
    if num is not None:
        nb, nf = add_textbox(slide, 12.3, 7.0, 0.9, 0.4)
        _set_run(nf.paragraphs[0].add_run(), str(num), 14, color=GRAY)
        nf.paragraphs[0].alignment = PP_ALIGN.RIGHT


def bullets(slide, items, left=0.7, top=1.55, width=11.9, size=20):
    height = min(5.3, 7.35 - top)
    tb, tf = add_textbox(slide, left, top, width, height)
    first = True
    for text, lvl in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lvl
        p.space_after = Pt(10 if lvl == 0 else 6)
        p.space_before = Pt(0)
        prefix = '\u2022  ' if lvl == 0 else '\u2013  '
        add_rich(p, prefix + text, size if lvl == 0 else size - 2)
    return tf


def figure(slide, name, caption, left, top, width):
    path = os.path.join(FIG, name)
    if not os.path.exists(path):
        print('ПРЕДУПРЕЖДЕНИЕ: нет рисунка', path)
        return
    w_px, h_px = Image.open(path).size
    height = width * h_px / w_px
    slide.shapes.add_picture(path, Inches(left), Inches(top),
                             width=Inches(width), height=Inches(height))
    tb, tf = add_textbox(slide, left, top + height + 0.05, width + 0.4, 0.5)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    add_rich(p, caption, 14, italic=True, color=BLACK)


def table(slide, rows, col_widths, left, top, width):
    n_r, n_c = len(rows), len(rows[0])
    shape = slide.shapes.add_table(n_r, n_c, Inches(left), Inches(top),
                                   Inches(width), Inches(0.55 * n_r))
    tbl = shape.table
    tbl.first_row = False
    tbl.horz_banding = False
    for j, w in enumerate(col_widths):
        tbl.columns[j].width = Inches(w)
    for i, row in enumerate(rows):
        for j, txt in enumerate(row):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PALE if i == 0 else WHITE
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.08); cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04); cell.margin_bottom = Inches(0.04)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            add_rich(p, txt, 16 if i == 0 else 15, bold=(i == 0))
    return tbl


def note(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    blank = prs.slide_layouts[6]

    # ---------------- 1. Титульный ----------------
    s = prs.slides.add_slide(blank)
    tb, tf = add_textbox(s, 0.9, 1.0, 11.5, 0.6)
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    _set_run(tf.paragraphs[0].add_run(), 'Конкурс «Энергетика: вызовы будущего»', 18, color=DARK)
    tb, tf = add_textbox(s, 0.7, 1.75, 12.0, 2.4, anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(),
             'Мультифрактальный анализ сигналов напряжения твердополимерного '
             'топливного элемента для диагностики деградации и неисправностей',
             30, bold=True)
    tb, tf = add_textbox(s, 0.9, 4.15, 11.5, 1.1)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(), 'Направление 2. Цифровые технологии, диагностика, мониторинг, '
                          'прогнозирование, моделирование и управление режимами '
                          'энергетических объектов', 16)
    tb, tf = add_textbox(s, 0.9, 5.35, 11.5, 1.9)
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(), 'Кондрашов Роман Андреевич', 20, bold=True)
    p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(), 'Уральский федеральный университет имени первого '
                          'Президента России Б.Н. Ельцина', 16)
    p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(), 'Физико-технологический институт, кафедра технической физики', 16)
    p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER
    _set_run(p.add_run(), '2026', 16)
    note(s, 'Название, ФИО и организация совпадают с титульным листом конкурсной работы.')

    # ---------------- 2. Проблема и актуальность ----------------
    s = prs.slides.add_slide(blank)
    title_bar(s, 'Проблема и актуальность', 2)
    bullets(s, [
        ('PEMFC — ключевой компонент водородной энергетики; барьер внедрения — ограниченный ресурс и деградация стека', 0),
        ('Эталонная диагностика EIS требует дорогого оборудования, дискретна, не работает онлайн', 0),
        ('Штатный сигнал напряжения используется лишь как среднее и СКО — информация о состоянии теряется', 0),
        ('Перемежаемые процессы (вода, тепло) делают сигнал мультифрактальным', 1),
    ], left=0.7, top=1.55, width=7.0, size=19)
    figure(s, 'fig01_degradation.png',
           'Рисунок 1 – Записи напряжений FC1/FC2', 7.9, 1.7, 5.0)
    note(s, 'Актуальность — до 10 баллов. FC1 — стационарный режим, FC2 — с пульсациями тока 5 кГц.')

    # ---------------- 3. Существующие решения ----------------
    s = prs.slides.add_slide(blank)
    title_bar(s, 'Существующие решения и их ограничения', 3)
    table(s, [
        ['Подход', 'Что даёт и в чём предел'],
        ['EIS + схемы R–L–CPE', 'Эталон состояния (R~ohm~, R~p~, α), но дискретно и дорого'],
        ['Статистика напряжения', 'Дёшево, онлайн, но низкая чувствительность к деградации'],
        ['Вейвлет-лидеры (FCLAB)', 'Диагностика отказов, но качественно, без связи с EIS, без нуль-контролей'],
        ['MFDFA / MFDCCA на ТЭ', 'Ранее не применялись (0 работ по данным поиска)'],
        ['Ниша', 'Количественная связь «мультифрактальные метрики ↔ EIS»'],
    ], [4.2, 7.7], 0.7, 1.5, 11.9)
    note(s, 'Сжатая версия таблицы 1 конкурсной работы.')

    # ---------------- 4. Цель и задачи ----------------
    s = prs.slides.add_slide(blank)
    title_bar(s, 'Цель и задачи', 4)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(1.5),
                             Inches(11.9), Inches(1.2))
    box.fill.solid(); box.fill.fore_color.rgb = PALE
    box.line.color.rgb = DARK; box.line.width = Pt(1.0)
    box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.25); tf.margin_right = Inches(0.25)
    add_rich(tf.paragraphs[0],
             'Цель: диагностика деградации и неисправностей PEMFC по мультифрактальным '
             'метрикам напряжения и количественная связь этих метрик с параметрами EIS',
             18)
    bullets(s, [
        ('Задачи', 0),
        ('1) реализовать и валидировать MFDFA, WLMA, MFDCCA', 1),
        ('2) полный анализ записей двух открытых датасетов (до 2680 ч)', 1),
        ('3) корреляции метрик с EIS — временная и поячеечная постановки', 1),
        ('4) отклик на «отказные» условия и алгоритм SOH-индикатора', 1),
    ], top=3.0, size=20)
    note(s, 'Задачи совпадают с разделом 2 конкурсной работы.')

    # ---------------- 5. Научная новизна ----------------
    s = prs.slides.add_slide(blank)
    title_bar(s, 'Научная новизна', 5)
    bullets(s, [
        ('Впервые: количественная связь мультифрактальных метрик напряжения с EIS — |ρ| до 0,95 (n = 8) и 0,89 (n = 10 ячеек)', 0),
        ('Первое сравнение MFDFA и WLMA на одних записях PEMFC с нуль-контролем и суррогатными тестами', 0),
        ('Измерено влияние пульсаций тока: Δh = 0,72 (FC2) против 0,50 (FC1)', 0),
        ('Открытый воспроизводимый протокол — в литературе по PEMFC отсутствовал', 0),
    ], left=0.7, top=1.55, width=7.0, size=18)
    figure(s, 'fig08_realhyfc_cells.png',
           'Рисунок 5 – Поячеечные Δh стека RealHyFC', 7.9, 1.7, 5.0)
    note(s, 'Новизна — до 25 баллов. Поячеечная валидация подтверждает временну́ю на независимом стеке.')

    # ---------------- 6. Предлагаемое решение ----------------
    s = prs.slides.add_slide(blank)
    title_bar(s, 'Предлагаемое решение', 6)
    bullets(s, [
        ('Аппарат: MFDFA, WLMA (вейвлет-лидеры), MFDCCA; метрики Δh, Δh~5~, h~0~, c~1~–c~3~', 0),
        ('Контроль качества: синтетика → кросс-валидация → нуль-контроль → суррогаты → диапазон масштабов', 0),
        ('Данные: IEEE PHM 2014 (1154/1020 ч) и RealHyFC (10 ячеек, 2678 ч); EIS: модель R–L–CPE', 0),
        ('Ключевое соотношение: F~q~(s) ∼ s^h(q)^,  τ(q) = q·h(q) − 1', 0),
    ], left=0.7, top=1.55, width=7.0, size=18)
    figure(s, 'fig07_validation.png',
           'Рисунок 2 – Валидация на синтетических эталонах', 7.9, 1.7, 5.0)
    note(s, 'Обоснованность — до 20 баллов. Погрешность MFDFA на каскадах ≤ 4,2 %.')

    # ---------------- 7. Результаты ----------------
    s = prs.slides.add_slide(blank)
    title_bar(s, 'Полученные результаты', 7)
    bullets(s, [
        ('Мультифрактальность реальна: нуль-контроль превышен в 4,3–28,2 раза', 0),
        ('Δh растёт с наработкой: ρ = 0,70, p = 2·10^-5^, +62 %', 0),
        ('Связь с EIS: Δh ↑ вместе с R~ohm~, R~p~; h~0~ ↔ α~CPE~; |ρ| до 0,95', 0),
        ('Голодание по воздуху: Δh~5~ 1,004 против 0,654; напряжение 0,672 В против 0,699 В', 0),
    ], left=0.7, top=1.5, width=11.9, size=18)
    figure(s, 'fig04_dh_trajectory_vs_EIS.png',
           'Рисунок 4 – Траектории Δh~5~ и R~p~ (FC1/FC2)', 0.9, 4.15, 4.6)
    note(s, 'Числа совпадают с разделом 5 конкурсной работы. FC2 (с пульсациями): тренда Δh нет, растёт h₀.')

    # ---------------- 8. Ожидаемый эффект ----------------
    s = prs.slides.add_slide(blank)
    title_bar(s, 'Ожидаемый эффект', 8)
    table(s, [
        ['Параметр', 'Значение'],
        ['Пользователь', 'Службы диагностики энергоустановок на ТЭ; разработчики СУ PEMFC'],
        ['Область применения', 'Бортовой мониторинг состояния стека; раннее предупреждение о деградации'],
        ['Технический эффект', 'Из штатного напряжения — информация, ранее доступная только EIS'],
        ['Экономический эффект', 'Меньше EIS-измерений и простоев; выше ресурс за счёт раннего вмешательства'],
    ], [3.6, 8.3], 0.7, 1.6, 11.9)
    note(s, 'Практическая значимость — до 25 баллов (таблица 6 конкурсной работы).')

    # ---------------- 9. Уровень готовности и внедрение ----------------
    s = prs.slides.add_slide(blank)
    title_bar(s, 'Уровень готовности и предложения по внедрению', 9)
    bullets(s, [
        ('Уровень готовности: расчётная модель, верифицированная на реальных открытых данных', 0),
        ('Что внедрять: бортовой SOH-индикатор по Δh (порог «пол + k·СКО», тренд по 5–10 окнам)', 0),
        ('Где: контроллеры энергоустановок на PEMFC; пересчёт окна — доли секунды (1–10 Гц в реальном времени)', 0),
    ], top=1.5, size=19)
    steps = [
        ('Напряжение\n1–10 Гц', '1'),
        ('MFDFA\n+ WLMA', '2'),
        ('Нуль-контроль\n«пол + k·СКО»', '3'),
        ('Тренд Δh\n→ SOH', '4'),
    ]
    x0, rw, aw = 0.7, 2.8, 0.28
    for k, (txt, num) in enumerate(steps):
        x = x0 + k * (rw + aw)
        box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.3),
                                 Inches(rw), Inches(1.15))
        box.fill.solid(); box.fill.fore_color.rgb = PALE
        box.line.color.rgb = DARK; box.line.width = Pt(1.0)
        box.shadow.inherit = False
        tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        lines = txt.split('\n')
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.CENTER
            add_rich(p, ln, 14, bold=(i == 0))
        if k < len(steps) - 1:
            ar = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,
                                    Inches(x + rw + 0.015), Inches(5.72),
                                    Inches(aw - 0.03), Inches(0.32))
            ar.fill.solid(); ar.fill.fore_color.rgb = DARK
            ar.line.fill.background()
            ar.shadow.inherit = False
    note(s, 'Шкала уровня готовности — п. 4.2.5 Положения. Следующий шаг — лабораторная проверка.')

    # ---------------- 10. Заключение ----------------
    s = prs.slides.add_slide(blank)
    title_bar(s, 'Заключение', 10)
    bullets(s, [
        ('Напряжение PEMFC мультифрактально; Δh — чувствительный индикатор деградации в квазистационарных режимах', 0),
        ('Впервые установлена количественная связь мультифрактальных метрик с параметрами EIS (|ρ| до 0,95)', 0),
        ('Методика готова к лабораторной проверке; мониторинг в реальном времени подтверждён расчётом', 0),
    ], top=1.5, size=20)
    box = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(5.3),
                             Inches(11.9), Inches(1.3))
    box.fill.solid(); box.fill.fore_color.rgb = PALE
    box.line.color.rgb = DARK; box.line.width = Pt(1.0)
    box.shadow.inherit = False
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    add_rich(p, 'Контакты: Кондрашов Роман Андреевич · roman200.kondrashov.2016@mail.ru · +7 982 621-10-69', 18)
    note(s, 'Выводы совпадают с заключением конкурсной работы.')

    prs.core_properties.title = 'Презентация конкурсной работы: мультифрактальный анализ напряжения PEMFC'
    prs.core_properties.author = 'Кондрашов Роман Андреевич'
    prs.save(OUT)
    print('saved', OUT, 'slides:', len(prs.slides._sldIdLst))


if __name__ == '__main__':
    build()
