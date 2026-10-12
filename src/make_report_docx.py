"""
make_docx.py -- convert docs/report.md into docs/report.docx (python-docx).

Handles headings (#..####), bullet/numbered lists, markdown tables and blockquotes.
"""
from __future__ import annotations

import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

PROJECT_ROOT = os.environ.get("PEMFC_PROJECT_ROOT") or os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), os.pardir))
DOCS = os.path.join(PROJECT_ROOT, "docs")
SRC = os.path.join(DOCS, "report.md")
DST = os.path.join(DOCS, "report.docx")


def clean(s: str) -> str:
    s = s.replace("**", "").replace("`", "")
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    return s


def main():
    with open(SRC, encoding="utf-8") as f:
        lines = f.read().split("\n")
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0)

    i = 0
    in_code = False
    while i < len(lines):
        ln = lines[i]
        if ln.strip().startswith("```"):
            in_code = not in_code
            i += 1
            continue
        if in_code:
            p = doc.add_paragraph(ln)
            p.style = doc.styles["No Spacing"]
            for r in p.runs:
                r.font.name = "Consolas"
                r.font.size = Pt(8)
            i += 1
            continue
        # table
        if ln.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").strip()) <= {"-", ":", " "}:
            header = [c.strip() for c in ln.strip().strip("|").split("|")]
            rows = []
            j = i + 2
            while j < len(lines) and lines[j].startswith("|"):
                rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                j += 1
            t = doc.add_table(rows=1, cols=len(header))
            t.style = "Light Grid Accent 1"
            for k, h in enumerate(header):
                t.rows[0].cells[k].text = clean(h)
                for r in t.rows[0].cells[k].paragraphs[0].runs:
                    r.font.bold = True
                    r.font.size = Pt(8)
            for row in rows:
                cells = t.add_row().cells
                for k in range(min(len(row), len(header))):
                    cells[k].text = clean(row[k])
                    for r in cells[k].paragraphs[0].runs:
                        r.font.size = Pt(8)
            doc.add_paragraph()
            i = j
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", ln)
        if m:
            lvl = len(m.group(1))
            doc.add_heading(clean(m.group(2)), level=min(lvl, 4))
            i += 1
            continue
        if ln.strip().startswith(">"):
            p = doc.add_paragraph(clean(ln.strip().lstrip(">").strip()))
            p.paragraph_format.left_indent = Cm(0.8)
            for r in p.runs:
                r.italic = True
            i += 1
            continue
        if re.match(r"^\s*[-*]\s+", ln):
            doc.add_paragraph(clean(re.sub(r"^\s*[-*]\s+", "", ln)), style="List Bullet")
            i += 1
            continue
        if re.match(r"^\s*\d+\.\s+", ln):
            doc.add_paragraph(clean(re.sub(r"^\s*\d+\.\s+", "", ln)), style="List Number")
            i += 1
            continue
        if ln.strip() == "---":
            i += 1
            continue
        if ln.strip():
            doc.add_paragraph(clean(ln))
        i += 1

    doc.save(DST)
    print("written", DST, os.path.getsize(DST), "bytes")


if __name__ == "__main__":
    main()
