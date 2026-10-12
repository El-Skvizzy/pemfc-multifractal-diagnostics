# -*- coding: utf-8 -*-
"""
make_article_docx.py — сборка конкурсной работы (DOCX) из docs/article.md
по правилам оформления из docs/Оформление_научной_статьи (измеренные параметры
исходного документа конкурса «Энергетика: вызовы будущего»).

Запуск:
  python src/make_article_docx.py ["$PEMFC_PROJECT_ROOT"]

Синтаксис article.md:
  [TITLEMETA] ... [ENDMETA]   — данные титульного листа
  [ABSTRACT] ...              — текст аннотации (один абзац)
  [KEYWORDS] ...              — абзац ключевых слов
  # Заголовок                 — заголовок раздела (Заголовок 1)
  #^ Заголовок                — заголовок раздела с новой страницы (приложение)
  ## Подраздел                — подраздел (14 пт полужирный, 6/4 пт)
  !! Вводная фраза            — абзац без красной строки
  - пункт списка              — маркированный список (маркер 0,63 / текст 1,27 см)
  [TCAPT] Таблица N           — номер таблицы; следующая строка — название
  | ... |                     — таблица (Сетка таблицы, 16 см)
  [W: a, b, c]                — ширины столбцов таблицы в см (после таблицы)
  [FIG N] файл | ширина_см    — рисунок по центру
  [CAP] Рисунок N – Название  — подпись рисунка
  [F] формула   (n)           — формула по центру, номер у правого края
  [REF] текст                 — пункт списка источников (нумерация автоматически)
  Инлайн: **полужирный**, ^верхний индекс^, ~нижний индекс~
"""
import os
import re
import sys

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TNR = 'Times New Roman'
PROJECT_ROOT = os.environ.get("PEMFC_PROJECT_ROOT") or os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), os.pardir))
ROOT = sys.argv[1] if len(sys.argv) > 1 else PROJECT_ROOT
MD = os.path.join(ROOT, 'docs', 'article.md')
FIGDIR = os.path.join(ROOT, 'results', 'figures')
if not os.path.isdir(FIGDIR):
    FIGDIR = os.path.join(ROOT, 'figures')
OUT = os.path.join(ROOT, 'docs', 'Кондрашов - Цифровые технологии - конкурсная работа.docx')


# ---------------------------------------------------------------- helpers
def set_font(style_or_run, name=TNR, size=None, bold=None, italic=None,
             color='000000'):
    f = style_or_run.font
    f.name = name
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = RGBColor.from_string(color)
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run.element, 'get_or_add_rPr') else None
    if rpr is not None:
        rf = rpr.find(qn('w:rFonts'))
        if rf is None:
            rf = OxmlElement('w:rFonts')
            rpr.insert(0, rf)
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
            rf.attrib.pop(qn(a), None)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
            rf.set(qn(a), name)


def fmt_par(par, before=None, after=None, line=None, first=None,
            left=None, align=None, keep_next=None, keep_lines=None):
    pf = par.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if first is not None:
        pf.first_line_indent = Cm(first)
    if left is not None:
        pf.left_indent = Cm(left)
    if align is not None:
        pf.alignment = align
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if keep_lines is not None:
        pf.keep_together = keep_lines
    return par


def ensure_style(doc, name, base='Normal'):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


INLINE = re.compile(r'(\*\*.+?\*\*|\^[^^]+?\^|~[^~]+?~)')


def _emit(par, text, size, bold, font, sup, sub):
    """Рекурсивный разбор инлайна: **жирный**, ^верхний^, ~нижний~ (в т. ч. вложенные)."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) > 4:
            _emit(par, part[2:-2], size, True, font, sup, sub)
        elif part.startswith('^') and part.endswith('^') and len(part) > 2:
            _emit(par, part[1:-1], size, bold, font, True, False)
        elif part.startswith('~') and part.endswith('~') and len(part) > 2:
            _emit(par, part[1:-1], size, bold, font, False, True)
        else:
            r = par.add_run(part)
            set_font(r, font, size, bold)
            if sup:
                r.font.superscript = True
            if sub:
                r.font.subscript = True


def add_inline(par, text, size=14, bold=False, font=TNR):
    _emit(par, text, size, bold, font, False, False)
    return par


def table_grid_xml(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    # фиксированная раскладка и ширина 16 см (9072 twips)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:type'), 'dxa')
    tblW.set(qn('w:w'), '9072')
    # все границы одинарные 0,5 пт
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + edge)
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:space'), '0')
        e.set(qn('w:color'), 'auto')
        borders.append(e)
    tblPr.append(borders)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader')
    trPr.append(h)


def cell_par(cell, text, size, bold):
    cell.text = ''
    p = cell.paragraphs[0]
    fmt_par(p, first=0, line=1.5, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_inline(p, text, size=size, bold=bold)


# ---------------------------------------------------------------- build
def build():
    md = open(MD, encoding='utf-8').read()

    # --- служебные блоки исключаются из тела документа
    md_body = re.sub(r'\[TITLEMETA\].*?\[ENDMETA\]\n?', '', md, flags=re.S)
    md_body = re.sub(r'\[ABSTRACT\].*?\[KEYWORDS\][^\n]*\n[^\n]*\n?', '',
                     md_body, flags=re.S)
    lines = md_body.split('\n')

    # --- разбор метаданных
    meta = {}
    m = re.search(r'\[TITLEMETA\](.*?)\[ENDMETA\]', md, re.S)
    if not m:
        sys.exit('нет блока TITLEMETA')
    for ln in m.group(1).strip().split('\n'):
        if ':' in ln:
            k, v = ln.split(':', 1)
            meta[k.strip()] = v.strip()

    m = re.search(r'\[ABSTRACT\]\s*(.*?)\n\[KEYWORDS\]', md, re.S)
    if not m:
        sys.exit('нет блока ABSTRACT/KEYWORDS')
    abstract = ' '.join(m.group(1).split())
    m2 = re.search(r'\[KEYWORDS\]\s*(.*?)\n\n#', md, re.S)
    keywords = ' '.join(m2.group(1).split()) if m2 else ''

    doc = Document()

    # --- страница
    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.left_margin = Cm(3.0); sec.right_margin = Cm(2.0)
    sec.top_margin = Cm(2.0); sec.bottom_margin = Cm(2.0)
    sec.header_distance = Cm(1.0); sec.footer_distance = Cm(1.0)

    # --- стили
    n = doc.styles['Normal']
    set_font(n, TNR, 14)
    fmt_par(n, before=0, after=0, line=1.5, first=1.25,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    h1 = doc.styles['Heading 1']
    set_font(h1, TNR, 16, bold=True)
    fmt_par(h1, before=0, after=6, line=1.5, first=0,
            align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True, keep_lines=True)

    h2 = ensure_style(doc, 'Podrazdel')
    set_font(h2, TNR, 14, bold=True)
    fmt_par(h2, before=6, after=4, line=1.5, first=0,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_next=True)

    bul = ensure_style(doc, 'Spisok')
    set_font(bul, TNR, 14)
    fmt_par(bul, before=0, after=0, line=1.5, first=-0.64, left=1.27,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    ref = ensure_style(doc, 'Istochnik')
    set_font(ref, TNR, 14)
    fmt_par(ref, before=0, after=6, line=1.5, first=0,
            align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # --- номера страниц: внизу справа, 14 пт, без номера на первой странице
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fmt_par(fp, before=6, line=1.0, first=0)
    run = fp.add_run()
    for tag, attrs in (('w:fldChar', {'w:fldCharType': 'begin'}),
                       ('w:instrText', {'xml:space': 'preserve'}),
                       ('w:fldChar', {'w:fldCharType': 'end'})):
        e = OxmlElement(tag)
        for a, v in attrs.items():
            e.set(qn(a), v)
        if tag == 'w:instrText':
            e.text = ' PAGE '
        run._r.append(e)
    set_font(run, TNR, 14)
    sec.different_first_page_header_footer = True

    # ---------------------------------------------------------------- титул
    def p0(align=WD_ALIGN_PARAGRAPH.LEFT):
        p = doc.add_paragraph()
        fmt_par(p, first=0, line=1.5, align=align)
        return p

    p = p0(WD_ALIGN_PARAGRAPH.CENTER); add_inline(p, meta['COMPETITION'], size=12)
    for _ in range(3):
        p0()
    p = p0(WD_ALIGN_PARAGRAPH.CENTER); add_inline(p, meta['TITLE'], size=16, bold=True)
    p0()
    p = p0(WD_ALIGN_PARAGRAPH.CENTER)
    mdir = re.match(r'^(Направление\s+\d+\.)(.*)$', meta['DIRECTION'])
    if mdir:
        add_inline(p, '**%s**%s' % (mdir.group(1), mdir.group(2)), size=12)
    else:
        add_inline(p, meta['DIRECTION'], size=12)
    p0(); p0()
    p = p0(WD_ALIGN_PARAGRAPH.RIGHT); add_inline(p, '**Автор:**', size=14)
    p = p0(WD_ALIGN_PARAGRAPH.RIGHT); add_inline(p, meta['AUTHOR_NAME'], size=14)
    p = p0(WD_ALIGN_PARAGRAPH.RIGHT); add_inline(p, meta['AUTHOR_STATUS'], size=14)
    p = p0(WD_ALIGN_PARAGRAPH.RIGHT); add_inline(p, meta['AUTHOR_PROGRAM'], size=14)
    p0(WD_ALIGN_PARAGRAPH.RIGHT)
    p = p0(WD_ALIGN_PARAGRAPH.RIGHT); add_inline(p, '**Организация:**', size=14)
    p = p0(WD_ALIGN_PARAGRAPH.RIGHT); add_inline(p, meta['ORG_NAME'], size=14)
    p = p0(WD_ALIGN_PARAGRAPH.RIGHT); add_inline(p, meta['ORG_DEPT'], size=14)
    # число пустых абзацев подобрано так, чтобы титул умещался на одну страницу
    for _ in range(3):
        p0(WD_ALIGN_PARAGRAPH.CENTER)
    p = p0(WD_ALIGN_PARAGRAPH.CENTER); add_inline(p, meta['CITY'], size=14)
    p = p0(WD_ALIGN_PARAGRAPH.CENTER); add_inline(p, meta['YEAR'], size=14)
    p = doc.add_paragraph(); fmt_par(p, first=0)
    p.add_run().add_break(WD_BREAK.PAGE)

    # ---------------------------------------------------------------- аннотация
    p = doc.add_paragraph()
    fmt_par(p, first=0, line=1.08, after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_inline(p, 'Аннотация', size=16, bold=False)
    for it in p.runs:
        it.font.italic = True
    p = doc.add_paragraph()
    fmt_par(p, first=1.25, line=1.08, after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_inline(p, abstract, size=14)
    p = doc.add_paragraph()
    fmt_par(p, first=1.25, line=1.08, after=8, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_inline(p, keywords, size=14)
    p = doc.add_paragraph(); fmt_par(p, first=0)
    p.add_run().add_break(WD_BREAK.PAGE)
    print('Аннотация: %d знаков (лимит 1500)' % len(abstract))

    # ---------------------------------------------------------------- тело
    refs = []
    ref_n = [0]
    fig_n = [0]
    tbl_n = [0]
    after_appendix = [False]
    text_cited = []

    def heading(text, page_break=False):
        p = doc.add_paragraph(style='Heading 1')
        if page_break:
            r = p.add_run(); r.add_break(WD_BREAK.PAGE)
        add_inline(p, text, size=16, bold=True)
        after_appendix[0] = False

    def subhead(text):
        p = doc.add_paragraph(style='Podrazdel')
        if after_appendix[0]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(p, text, size=14, bold=True)

    def para(text, first=1.25):
        p = doc.add_paragraph()
        fmt_par(p, first=first, line=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        add_inline(p, text, size=14)
        text_cited.append(text)
        return p

    def bullet(text):
        p = doc.add_paragraph(style='Spisok')
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.27), WD_TAB_ALIGNMENT.LEFT)
        r = p.add_run('\u2022\t'); set_font(r, 'Symbol', 14)
        add_inline(p, text, size=14)
        text_cited.append(text)

    def table_caption(label, name):
        p = doc.add_paragraph()
        fmt_par(p, first=0, line=1.5, align=WD_ALIGN_PARAGRAPH.RIGHT)
        add_inline(p, label, size=10)
        p = doc.add_paragraph()
        fmt_par(p, first=0, line=1.5, align=WD_ALIGN_PARAGRAPH.CENTER,
                keep_next=True)
        add_inline(p, name, size=12)

    def render_table(header, rows, widths):
        ncols = max(len(header), max((len(r) for r in rows), default=0))
        if widths is None:
            widths = [16.0 / ncols] * ncols
        s = sum(widths)
        widths = [w * 16.0 / s for w in widths]
        t = doc.add_table(rows=1, cols=ncols)
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        table_grid_xml(t)
        repeat_header(t.rows[0])
        for j in range(ncols):
            cell = t.rows[0].cells[j]
            cell.width = Cm(widths[j])
            cell_par(cell, header[j] if j < len(header) else '', 12, True)
        for row in rows:
            tr = t.add_row()
            for j in range(ncols):
                cell = tr.cells[j]
                cell.width = Cm(widths[j])
                cell_par(cell, row[j] if j < len(row) else '', 11, False)
        gap = doc.add_paragraph()
        fmt_par(gap, first=0, line=1.0)
        r = gap.add_run(); set_font(r, TNR, 4)

    def figure(fname, width, caption):
        fig_n[0] += 1
        path = os.path.join(FIGDIR, fname)
        if not os.path.exists(path):
            print('ПРЕДУПРЕЖДЕНИЕ: нет файла рисунка', path)
        p = doc.add_paragraph()
        fmt_par(p, first=0, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
        p.add_run().add_picture(path, width=Cm(width))
        p = doc.add_paragraph()
        fmt_par(p, first=0, line=1.0, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
        add_inline(p, caption, size=12)

    def formula(text):
        mf = re.match(r'^(.*?)\s+\((\d+)\)\s*$', text)
        body, num = (mf.group(1), mf.group(2)) if mf else (text, None)
        parts = [s.strip() for s in body.split('||') if s.strip()]
        for k, part in enumerate(parts):
            last = k == len(parts) - 1
            p = doc.add_paragraph()
            fmt_par(p, first=0, before=(12 if k == 0 else 0),
                    after=(12 if last else 0), align=WD_ALIGN_PARAGRAPH.CENTER)
            if last and num is not None:
                p.paragraph_format.tab_stops.add_tab_stop(Cm(8.0), WD_TAB_ALIGNMENT.CENTER)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(16.0), WD_TAB_ALIGNMENT.RIGHT)
                r = p.add_run('\t'); set_font(r, 'Cambria Math', 14)
                add_inline(p, part, size=14, font='Cambria Math')
                r = p.add_run('\t(%s)' % num); set_font(r, TNR, 14)
            else:
                add_inline(p, part, size=14, font='Cambria Math')

    def reference(text):
        ref_n[0] += 1
        p = doc.add_paragraph(style='Istochnik')
        add_inline(p, '%d. %s' % (ref_n[0], text), size=14)
        refs.append(text)

    i = 0
    while i < len(lines):
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith('[TITLEMETA]') or s.startswith('[ABSTRACT]') \
                or s.startswith('[KEYWORDS]'):
            i += 1
            continue
        if s.startswith('[ENDMETA]'):
            i += 1
            continue
        if s.startswith('#^ '):
            heading(s[3:].strip(), page_break=True)
            after_appendix[0] = True
            i += 1
            continue
        if s.startswith('# '):
            heading(s[2:].strip())
            i += 1
            continue
        if s.startswith('## '):
            subhead(s[3:].strip())
            i += 1
            continue
        if s.startswith('!! '):
            para(s[3:].strip(), first=0)
            i += 1
            continue
        if s.startswith('- '):
            buf = []
            while i < len(lines) and lines[i].strip().startswith('- '):
                buf.append(lines[i].strip()[2:])
                i += 1
            for it in buf:
                bullet(it)
            continue
        if s.startswith('[TCAPT] '):
            label = s[len('[TCAPT] '):].strip()
            name_line = lines[i + 1].strip() if i + 1 < len(lines) else ''
            table_caption(label, name_line)
            i += 2
            continue
        if s.startswith('|'):
            def split_row(line):
                t = line.strip()
                if t.startswith('|'):
                    t = t[1:]
                if t.endswith('|'):
                    t = t[:-1]
                return [c.strip().replace('\\|', '|')
                        for c in re.split(r'(?<!\\)\|', t)]
            header = split_row(s)
            i += 1
            # строка-разделитель markdown (|---|---|) может отсутствовать —
            # пропускаем её только если она действительно есть
            if i < len(lines) and re.match(r'^\|[\s\-:|]+\|$', lines[i].strip()):
                i += 1
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                rows.append(split_row(lines[i]))
                i += 1
            widths = None
            if i < len(lines) and lines[i].strip().startswith('[W:'):
                widths = [float(x) for x in
                          re.findall(r'[\d.]+', lines[i])]
                i += 1
            render_table(header, rows, widths)
            continue
        if s.startswith('[FIG '):
            mf = re.match(r'\[FIG \d+\]\s*(\S+?)\s*\|\s*([\d.]+)', s)
            fname, width = mf.group(1), float(mf.group(2))
            cap = ''
            if i + 1 < len(lines) and lines[i + 1].strip().startswith('[CAP]'):
                cap = lines[i + 1].strip()[len('[CAP] '):].strip()
                i += 1
            figure(fname, width, cap)
            i += 1
            continue
        if s.startswith('[F] '):
            formula(s[len('[F] '):].strip())
            i += 1
            continue
        if s.startswith('[REF] '):
            reference(s[len('[REF] '):].strip())
            i += 1
            continue
        # обычный абзац: собирать строки до пустой строки/директивы
        buf = [s]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or re.match(r'^(#|##|!!|- |\[|$)', nxt):
                break
            if nxt.startswith('|'):
                break
            buf.append(nxt)
            i += 1
        para(' '.join(buf))

    doc.save(OUT)

    # --- контроль ссылок
    cited = set()
    for t in text_cited:
        for m in re.finditer(r'\[([\d,\s]+)\]', t):
            for x in m.group(1).split(','):
                if x.strip().isdigit():
                    cited.add(int(x.strip()))
    if cited:
        mx = max(cited)
        missing = [k for k in range(1, mx + 1) if k not in cited]
        print('Ссылки в тексте: %d..%d (всего %d)' % (min(cited), mx, len(cited)))
        if missing:
            print('НЕ УПОМЯНУТЫ в тексте:', missing)
    print('Записей в списке источников:', len(refs))
    print('Соответствие: макс. номер ссылки %d, записей %d %s'
          % ((max(cited) if cited else 0), len(refs),
             'OK' if (cited and max(cited) == len(refs) and
                      all(k in cited for k in range(1, len(refs) + 1)))
             else 'ПРОВЕРИТЬ'))
    print('Сохранено:', OUT)


if __name__ == '__main__':
    build()
